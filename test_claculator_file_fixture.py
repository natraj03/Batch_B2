import openpyxl
import pytest

from  calculator_with_file import CalculatorFromFile


workbook = None

def setup_function():
    global workbook
    workbook = openpyxl.load_workbook("practicefile.xlsx")

def teardown_function():
    global workbook
    if workbook:
        workbook.close()

#
# def test_mul():
#     current_sheet = workbook.active
#     values = []
#     for row in current_sheet.iter_rows(values_only=True):
#         print("row",row)
#         values.append(row)
#     # for colm in current_sheet.iter_cols(values_only=True):
#     #     print("column", colm)
#
#     print("values in mul",values)
#     print("values in mul at index 1",values[1])
#     mutobj = CalculatorFromFile()
#
#     row0 = values[0]
#     print(mutobj.mult(row0[0],row0[1]))
#     row1 = values[1]
#     print(mutobj.mult(row1[0], row1[1]))
#     row2 = values[2]
#     print(mutobj.mult(row2[0], row2[1]))
#     row3 = values[3]
#     print(mutobj.mult(row3[0], row3[1]))
#
# #
# def test_add():
#     current_sheet = workbook.active
#     values = []
#     for row in current_sheet.iter_rows(values_only=True):
#         values.append(row)
#     print("values in add", values)
#     mutobj = CalculatorFromFile()
#
#     row0 = values[0]
#     print(mutobj.add(row0[0], row0[1]))
#     row1 = values[1]
#     print(mutobj.add(row1[0], row1[1]))
#     row2 = values[2]
#     print(mutobj.add(row2[0], row2[1]))
#     row3 = values[3]
#     print(mutobj.add(row3[0], row3[1]))
# def test_sub():
#     current_sheet = workbook.active
#     values = []
#     for row in current_sheet.iter_rows(values_only=True):
#         values.append(row)
#     print("values in sun", values)
#     mutobj = CalculatorFromFile()
#
#     row0 = values[0]
#     print(mutobj.sub(row0[0], row0[1]))
#     row1 = values[1]
#     print(mutobj.sub(row1[0], row1[1]))
#     row2 = values[2]
#     print(mutobj.sub(row2[0], row2[1]))
#     row3 = values[3]
#     print(mutobj.sub(row3[0], row3[1]))
#

############  using fixture ######################
@pytest.fixture
def all_values():
    current_sheet = workbook.active
    values = []
    for row in current_sheet.iter_rows(values_only=True):
        values.append(row)
    return  values

    # all_values = all_values()
def test_add_preload_data(all_values):
    print("values in add", all_values)
    mutobj = CalculatorFromFile()

    row0 = all_values[0]
    assert mutobj.add(row0[0], row0[1]) == 15.0
    row1 = all_values[1]
    print(mutobj.add(row1[0], row1[1]))
    row2 = all_values[2]
    print(mutobj.add(row2[0], row2[1]))
    row3 = all_values[3]
    print(mutobj.add(row3[0], row3[1]))

def test_sub_preload_data(all_values):

    print("values in sun", all_values)
    mutobj = CalculatorFromFile()

    row0 = all_values[0]
    assert mutobj.sub(row0[0], row0[1]) == -5.0
    row1 = all_values[1]
    print(mutobj.sub(row1[0], row1[1]))
    row2 = all_values[2]
    print(mutobj.sub(row2[0], row2[1]))
    row3 = all_values[3]
    print(mutobj.sub(row3[0], row3[1]))

def test_mul_preload_data(all_values):
    print("values in mul",all_values)
    mutobj = CalculatorFromFile()

    row0 = all_values[0]
    assert mutobj.mult(row0[0],row0[1]) == 50.0
    row1 = all_values[1]
    print(mutobj.mult(row1[0], row1[1]))
    row2 = all_values[2]
    print(mutobj.mult(row2[0], row2[1]))
    row3 = all_values[3]
    print(mutobj.mult(row3[0], row3[1]))
