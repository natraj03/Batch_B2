import pytest
import openpyxl

from  Revision_python_pytest import *


excelFile = None

def setup_function():
    global  excelFile
    excelFile = openpyxl.load_workbook("input_file.xlsx")

def teardown_function():
    excelFile.close()


def test_multiply():
    sheet1 = excelFile['Multiplication']
    calObj = Calculator()
    for row in sheet1.iter_rows(values_only=True):
        if type(row[0]) == int and type(row[1]) == int:
            temp_result = calObj.multiply(row[0], row[1])
            print("input 1=",row[0],"input2",row[1],"Result",temp_result,"expected rresult",row[2])
            assert int(temp_result) == row[2]

def test_add():
    sheet1 = excelFile['Addition']
    calObj = Calculator()
    for row in sheet1.iter_rows(values_only=True):
        temp_result = calObj.add(row[0], row[1])
        print("input 1=",row[0],"input2",row[1],"Result",temp_result,"expected rresult",row[2])
        assert int(temp_result) == row[2]


def test_plendrome():
    sheet1 = excelFile['Palendrome']
    palenObj = PalendromeStr()
    for row in sheet1.iter_rows(values_only=True):
        temp_result = palenObj.ispalendrome(str(row[0]))
        assert int(temp_result) == int(row[1])

def test_usingLambda():
    sheet1 = excelFile['Multiplication']
    for row in sheet1.iter_rows(values_only=True):
        if type(row[0]) == int and type(row[1]) == int:
            x = row[0]
            y = row[1]
            # assert lambda x, y: x * y == row[2]
            # assert x * y == row[2]

